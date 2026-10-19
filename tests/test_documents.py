from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document

from talent_agent.documents import (
    CVEditorAgent,
    CoverLetterWriterAgent,
    generate_application_docs,
)


class ApplicationDocumentTests(unittest.TestCase):
    def test_agent_prompts_use_role_responsibilities_and_candidate_evidence(self) -> None:
        job = {
            "title": "Platform Engineer",
            "company": "Example Ltd",
            "location": "London",
            "employment_type": "permanent",
            "description": "Own Kubernetes platform reliability and improve observability.",
        }
        cv_agent = CVEditorAgent.__new__(CVEditorAgent)
        cv_agent.agent = Mock()
        cv_agent.agent.complete.return_value = "CV"
        letter_agent = CoverLetterWriterAgent.__new__(CoverLetterWriterAgent)
        letter_agent.agent = Mock()
        letter_agent.agent.complete.return_value = "Letter"

        cv_agent.create("Candidate operated Kubernetes clusters.", job)
        letter_agent.create("Candidate operated Kubernetes clusters.", job, {})

        cv_system, cv_prompt = cv_agent.agent.complete.call_args.args
        _, letter_prompt = letter_agent.agent.complete.call_args.args
        for prompt in (cv_prompt, letter_prompt):
            self.assertIn("Own Kubernetes platform reliability", prompt)
            self.assertIn("improve observability", prompt)
            self.assertIn("Candidate operated Kubernetes clusters.", prompt)
        self.assertIn("not merely its title", cv_system)
        self.assertIn("not just the role title", letter_prompt)
        self.assertIn("two or three important role needs", letter_prompt)

    def test_documents_omit_template_labels_and_include_github_url(self) -> None:
        github_url = "https://github.com/example-user"
        cv_agent = Mock()
        cv_agent.create.return_value = "# Nadeem Bhatti\n\n# Tailored CV\n\nPlatform engineer\n\n[Date]"
        letter_agent = Mock()
        letter_agent.create.return_value = (
            "# Cover Letter — Example Ltd\n\nDear Hiring Manager,\n\n"
            "I can contribute to the platform team.\n\nKind regards,\nNadeem\n\n[Date]"
        )
        job = {
            "job_id": "role-1",
            "company": "Example Ltd",
            "title": "Platform Engineer",
        }

        with tempfile.TemporaryDirectory() as temporary_dir:
            directory = Path(temporary_dir)
            with (
                patch("talent_agent.documents.CVEditorAgent", return_value=cv_agent),
                patch("talent_agent.documents.CoverLetterWriterAgent", return_value=letter_agent),
            ):
                documents = generate_application_docs(
                    "Master CV", job, {}, directory, output_format="docx", github_url=github_url
                )

            cv_content = "\n".join(paragraph.text for paragraph in Document(directory / documents["cv_document"]).paragraphs)
            letter_content = "\n".join(
                paragraph.text for paragraph in Document(directory / documents["cover_letter_document"]).paragraphs
            )

        for content in (cv_content, letter_content):
            self.assertNotIn("Tailored CV", content)
            self.assertNotIn("Cover Letter", content)
            self.assertNotIn("[Date]", content)
            self.assertIn(github_url, content)

        cv_agent.create.assert_called_once_with("Master CV", job, github_url)
        letter_agent.create.assert_called_once_with("Master CV", job, {}, github_url)

    def test_safety_metadata_only_checkpoint_is_not_written_as_a_document(self) -> None:
        cv_agent = Mock()
        cv_agent.create.return_value = "Nadeem Bhatti\nPlatform engineer"
        letter_agent = Mock()
        letter_agent.create.return_value = "Dear Hiring Manager,\n\nI am applying for this role."
        job = {"job_id": "role-1", "company": "Example Ltd", "title": "Platform Engineer"}

        with tempfile.TemporaryDirectory() as temporary_dir:
            directory = Path(temporary_dir)
            with (
                patch("talent_agent.documents.CVEditorAgent", return_value=cv_agent),
                patch("talent_agent.documents.CoverLetterWriterAgent", return_value=letter_agent),
            ):
                documents = generate_application_docs(
                    "Master CV",
                    job,
                    {},
                    directory,
                    output_format="docx",
                    cv_text="User Safety: safe",
                    cover_letter_text="User Safety: safe",
                )

            cv_content = "\n".join(
                paragraph.text for paragraph in Document(directory / documents["cv_document"]).paragraphs
            )
            letter_content = "\n".join(
                paragraph.text for paragraph in Document(directory / documents["cover_letter_document"]).paragraphs
            )

        self.assertIn("Platform engineer", cv_content)
        self.assertIn("Dear Hiring Manager", letter_content)
        cv_agent.create.assert_called_once()
        letter_agent.create.assert_called_once()

    def test_model_commentary_is_rejected_and_regenerated(self) -> None:
        cv_agent = Mock()
        cv_agent.create.side_effect = [
            "The user wants me to create a CV.\nLet me analyze the role.",
            "Nadeem Bhatti\nPlatform engineer with verified Kubernetes experience.",
        ]
        letter_agent = Mock()
        letter_agent.create.return_value = "Dear Hiring Manager,\n\nI built Kubernetes platforms."
        job = {"job_id": "role-1", "company": "Example Ltd", "title": "Platform Engineer"}

        with tempfile.TemporaryDirectory() as temporary_dir:
            with (
                patch("talent_agent.documents.CVEditorAgent", return_value=cv_agent),
                patch("talent_agent.documents.CoverLetterWriterAgent", return_value=letter_agent),
            ):
                result = generate_application_docs(
                    "Master CV", job, {}, Path(temporary_dir), output_format="md"
                )

            cv_content = (Path(temporary_dir) / result["cv_document"]).read_text()

        self.assertTrue(cv_content.startswith("Nadeem Bhatti"))
        self.assertNotIn("Let me analyze", cv_content)
        self.assertEqual(cv_agent.create.call_count, 2)
        self.assertIn("no analysis", cv_agent.create.call_args.args[3][-1])

    def test_unapproved_review_does_not_write_application_documents(self) -> None:
        cv_agent = Mock()
        cv_agent.create.return_value = "Nadeem Bhatti\nPlatform engineer."
        letter_agent = Mock()
        letter_agent.create.return_value = "Dear Hiring Manager,\n\nI built platforms."
        reviewer = Mock()
        reviewer.review.return_value = {
            "approved": False,
            "cv_findings": [],
            "cover_letter_findings": [],
            "summary": "Unsupported claims remain.",
        }
        job = {"job_id": "role-1", "company": "Example Ltd", "title": "Platform Engineer"}
        progress: list[dict[str, object]] = []

        with tempfile.TemporaryDirectory() as temporary_dir:
            with (
                patch("talent_agent.documents.CVEditorAgent", return_value=cv_agent),
                patch("talent_agent.documents.CoverLetterWriterAgent", return_value=letter_agent),
            ):
                with self.assertRaisesRegex(ValueError, "did not approve.*Unsupported claims remain"):
                    generate_application_docs(
                        "Master CV",
                        job,
                        {},
                        Path(temporary_dir),
                        output_format="md",
                        reviewer=reviewer,
                        on_review_progress=lambda record: progress.append(dict(record)),
                    )

            self.assertEqual(list(Path(temporary_dir).glob("*.md")), [])

        self.assertEqual(reviewer.review.call_count, 2)
        self.assertEqual(progress[-1]["status"], "revision_pending")

    def test_reviewer_feedback_triggers_one_revision_and_a_final_review(self) -> None:
        cv_agent = Mock()
        cv_agent.create.side_effect = ["Initial CV", "Revised CV"]
        letter_agent = Mock()
        letter_agent.create.side_effect = ["Initial letter", "Revised letter"]
        reviewer = Mock()
        reviewer.review.side_effect = [
            {
                "approved": False,
                "cv_findings": ["CV lacks evidence for the role's reliability responsibility."],
                "cover_letter_findings": ["Letter does not name the hiring company."],
                "summary": "Role alignment needs improvement.",
            },
            {
                "approved": True,
                "cv_findings": [],
                "cover_letter_findings": [],
                "summary": "Both documents passed.",
            },
        ]
        review_progress: list[dict[str, object]] = []
        job = {
            "job_id": "role-1",
            "company": "Example Ltd",
            "title": "Platform Engineer",
            "description": "Improve platform reliability.",
        }

        with tempfile.TemporaryDirectory() as temporary_dir:
            with (
                patch("talent_agent.documents.CVEditorAgent", return_value=cv_agent),
                patch("talent_agent.documents.CoverLetterWriterAgent", return_value=letter_agent),
            ):
                result = generate_application_docs(
                    "Master CV",
                    job,
                    {},
                    Path(temporary_dir),
                    output_format="md",
                    reviewer=reviewer,
                    linkedin_profile={"status": "available", "text": "Candidate profile"},
                    on_review_progress=lambda record: review_progress.append(dict(record)),
                )

            self.assertEqual((Path(temporary_dir) / result["cv_document"]).read_text().strip(), "Revised CV")
            self.assertEqual(
                (Path(temporary_dir) / result["cover_letter_document"]).read_text().strip(),
                "Revised letter",
            )

        self.assertTrue(result["document_review"]["approved"])
        self.assertEqual(reviewer.review.call_count, 2)
        self.assertEqual(cv_agent.create.call_count, 2)
        self.assertIn("reliability responsibility", cv_agent.create.call_args.args[3][0])
        self.assertIn("hiring company", letter_agent.create.call_args.args[4][0])
        self.assertTrue(any(record["status"] == "revision_pending" for record in review_progress))
        self.assertEqual(review_progress[-1]["status"], "complete")


if __name__ == "__main__":
    unittest.main()
